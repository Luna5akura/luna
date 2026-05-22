import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("src/posts/Social Research")
AUTHOR = "王鲲淏-22307100011"


def strip_comments(text: str) -> str:
    lines = []
    for line in text.splitlines():
        if "%" in line:
            pos = line.find("%")
            if pos == 0 or line[pos - 1] != "\\":
                line = line[:pos]
        lines.append(line)
    return "\n".join(lines)


def extract_macro(text: str, name: str) -> str:
    pattern = re.compile(rf"\\{name}\{{(.*?)\}}", re.DOTALL)
    match = pattern.search(text)
    return match.group(1).strip() if match else ""


def clean_inline(text: str) -> str:
    text = text.replace("\\&", "&").replace("\\%", "%").replace("\\$", "$")
    text = text.replace("\\_", "_")
    text = text.replace("\\rightarrow", "->").replace("\\times", "x")
    text = text.replace("\\alpha", "alpha").replace("\\beta", "beta")
    text = text.replace("\\gamma", "gamma")
    text = text.replace("\\LaTeX", "LaTeX")
    text = text.replace("\\\\", "\n")
    text = re.sub(r"\$(.*?)\$", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\textit\{(.*?)\}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\alert\{(.*?)\}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\textbf\{(.*?)\}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\emph\{(.*?)\}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\bfseries", "", text)
    text = re.sub(r"\\(Large|large|small|scriptsize|tiny|normalsize|huge|Huge|centering)\b", "", text)
    text = re.sub(r"\\vspace\{.*?\}", "", text)
    text = re.sub(r"\\hspace\{.*?\}", "", text)
    text = re.sub(r"\\begin\{.*?\}|\\end\{.*?\}", "", text)
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?(?:\{[^{}]*\})?", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def extract_environment_blocks(text: str, env: str):
    pattern = re.compile(rf"\\begin\{{{env}\}}(?:\[[^\]]*\])?(.*?)\\end\{{{env}\}}", re.DOTALL)
    return list(pattern.finditer(text))


def parse_items(block_text: str, numbered: bool = False):
    parts = re.split(r"\\item", block_text)
    items = []
    index = 1
    for part in parts[1:]:
        item = clean_inline(part)
        if not item:
            continue
        prefix = f"{index}. " if numbered else "• "
        items.append(prefix + item)
        index += 1
    return items


def parse_table(block_text: str):
    rows = []
    for raw_row in re.split(r"\\\\", block_text):
        row = raw_row.strip()
        if not row:
            continue
        if any(token in row for token in ["\\toprule", "\\midrule", "\\bottomrule"]):
            continue
        cells = [clean_inline(cell) for cell in row.split("&")]
        cells = [cell for cell in cells if cell]
        if cells:
            rows.append(" | ".join(cells))
    return rows


def parse_tikz(block_text: str):
    lines = []
    node_texts = re.findall(
        r"\\node(?:\[[^\]]*\])?(?:\s*\([^)]+\))?(?:,\s*[^\{]+)?\s*\{([^{}]+)\};",
        block_text,
        flags=re.DOTALL,
    )
    label_texts = re.findall(r"node\[[^\]]*\]\s*\{([^{}]+)\}", block_text, flags=re.DOTALL)
    for text in node_texts + label_texts:
        cleaned = clean_inline(text)
        if cleaned and cleaned not in lines:
            lines.append(cleaned.replace("\n", " "))
    return lines


def extract_plain_lines(text: str):
    cleaned = clean_inline(text)
    lines = []
    for line in cleaned.splitlines():
        line = line.strip()
        if line:
            lines.append(line)
    return lines


def frame_to_lines(body: str):
    work = body
    lines = []

    title_match = re.search(r"\\bigidea\{(.*?)\}", work, re.DOTALL)
    title = clean_inline(title_match.group(1)) if title_match else ""
    work = re.sub(r"\\bigidea\{.*?\}", "", work, flags=re.DOTALL)

    for match in extract_environment_blocks(work, "itemize"):
        lines.extend(parse_items(match.group(1), numbered=False))
    work = re.sub(r"\\begin\{itemize\}(.*?)\\end\{itemize\}", "", work, flags=re.DOTALL)

    for match in extract_environment_blocks(work, "enumerate"):
        lines.extend(parse_items(match.group(1), numbered=True))
    work = re.sub(r"\\begin\{enumerate\}(.*?)\\end\{enumerate\}", "", work, flags=re.DOTALL)

    for match in extract_environment_blocks(work, "tabular"):
        table_lines = parse_table(match.group(1))
        if table_lines:
            if lines:
                lines.append("")
            lines.extend(table_lines)
    work = re.sub(r"\\begin\{tabular\}(.*?)\\end\{tabular\}", "", work, flags=re.DOTALL)

    for match in extract_environment_blocks(work, "tikzpicture"):
        tikz_lines = parse_tikz(match.group(1))
        if tikz_lines:
            if lines:
                lines.append("")
            lines.extend("• " + line for line in tikz_lines)
    work = re.sub(r"\\begin\{tikzpicture\}(.*?)\\end\{tikzpicture\}", "", work, flags=re.DOTALL)

    plain_lines = extract_plain_lines(work)
    if plain_lines:
        if lines:
            lines.append("")
        lines.extend(plain_lines)

    compact = [line for line in lines if line or (lines and line == "")]
    return title, compact


def add_textbox(slide, left, top, width, height, text, font_size=24, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return tf


def add_multiline_content(slide, lines):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    non_empty_count = len([line for line in lines if line.strip()])
    if non_empty_count <= 6:
        size = 24
    elif non_empty_count <= 10:
        size = 20
    elif non_empty_count <= 16:
        size = 18
    else:
        size = 14

    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line if line else " "
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        if line.startswith("• ") or re.match(r"\d+\. ", line):
            p.level = 0
    return tf


def parse_document(text: str):
    text = strip_comments(text)
    title = clean_inline(extract_macro(text, "title"))
    author = clean_inline(extract_macro(text, "author"))
    date = clean_inline(extract_macro(text, "date"))

    body_match = re.search(r"\\begin\{document\}(.*)\\end\{document\}", text, re.DOTALL)
    body = body_match.group(1) if body_match else ""

    slides = []
    idx = 0
    while idx < len(body):
        if body.startswith("\\begin{frame}[plain]", idx):
            end = body.find("\\end{frame}", idx)
            chunk = body[idx:end + len("\\end{frame}")]
            if "\\titlepage" in chunk:
                slides.append({"type": "title", "title": title, "author": author, "date": date})
            else:
                inner = re.search(r"\\begin\{frame\}(?:\[[^\]]*\])?(.*)\\end\{frame\}", chunk, re.DOTALL)
                frame_body = inner.group(1) if inner else chunk
                frame_title, lines = frame_to_lines(frame_body)
                slides.append({"type": "content", "title": frame_title, "lines": lines})
            idx = end + len("\\end{frame}")
            continue

        if body.startswith("\\sectionpageframe{", idx):
            start = idx + len("\\sectionpageframe{")
            end = body.find("}", start)
            section_title = clean_inline(body[start:end])
            slides.append({"type": "section", "title": section_title})
            idx = end + 1
            continue

        if body.startswith("\\begin{frame}", idx):
            end = body.find("\\end{frame}", idx)
            chunk = body[idx:end + len("\\end{frame}")]
            inner = re.search(r"\\begin\{frame\}(?:\[[^\]]*\])?(.*)\\end\{frame\}", chunk, re.DOTALL)
            frame_body = inner.group(1) if inner else chunk
            frame_title, lines = frame_to_lines(frame_body)
            slides.append({"type": "content", "title": frame_title, "lines": lines})
            idx = end + len("\\end{frame}")
            continue

        idx += 1

    return title, author, date, slides


def build_pptx(tex_path: Path):
    text = tex_path.read_text(encoding="utf-8")
    title, author, date, slides = parse_document(text)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    for slide_info in slides:
        slide = prs.slides.add_slide(blank)
        if slide_info["type"] == "title":
            add_textbox(slide, Inches(1.0), Inches(1.4), Inches(11.3), Inches(1.8), title, font_size=28, bold=True, align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(2.0), Inches(3.2), Inches(9.3), Inches(0.7), author, font_size=20, align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(2.0), Inches(4.0), Inches(9.3), Inches(0.7), date, font_size=16, align=PP_ALIGN.CENTER)
        elif slide_info["type"] == "section":
            add_textbox(slide, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.2), slide_info["title"], font_size=30, bold=True, align=PP_ALIGN.CENTER)
        else:
            title_text = slide_info["title"] or title or tex_path.stem
            add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8), title_text, font_size=24, bold=True, align=PP_ALIGN.LEFT)
            add_multiline_content(slide, slide_info["lines"])

    out_path = tex_path.with_suffix(".pptx")
    prs.save(out_path)
    return len(slides), out_path


def main():
    tex_files = sorted(ROOT.glob("*.tex"))
    for tex_path in tex_files:
        slides, out_path = build_pptx(tex_path)
        print(f"{tex_path.name} -> {out_path.name} ({slides} slides)")


if __name__ == "__main__":
    main()
